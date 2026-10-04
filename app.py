import os
import sqlite3
from datetime import datetime
from functools import wraps
from io import BytesIO
from flask import Flask, render_template, request, redirect, url_for, session, send_file

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

app = Flask(__name__)
app.secret_key = 'clave_secreta_universidad_inventario'

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATABASE = os.path.join(BASE_DIR, 'database.db')

def get_db():
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    with get_db() as db:
        db.execute('''
            CREATE TABLE IF NOT EXISTS usuarios (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                password TEXT NOT NULL,
                rol TEXT NOT NULL
            )
        ''')

        db.execute('''
            CREATE TABLE IF NOT EXISTS secciones (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nombre TEXT NOT NULL,
                tipo TEXT NOT NULL
            )
        ''')

        # --- NUEVO: Categorías de equipos (Ej: Cómputo, Red, Herramienta) ---
        db.execute('''
            CREATE TABLE IF NOT EXISTS categorias (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nombre TEXT UNIQUE NOT NULL
            )
        ''')

        # --- NUEVO: Códigos/prefijos predefinidos por categoría (Ej: MON, CPU) ---
        # 'contador' guarda el último número usado para autonumerar (MON-001, MON-002...)
        db.execute('''
            CREATE TABLE IF NOT EXISTS codigos (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                categoria_id INTEGER NOT NULL,
                prefijo TEXT UNIQUE NOT NULL,
                descripcion TEXT,
                contador INTEGER NOT NULL DEFAULT 0,
                FOREIGN KEY (categoria_id) REFERENCES categorias (id)
            )
        ''')

        db.execute('''
            CREATE TABLE IF NOT EXISTS elementos (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                codigo TEXT UNIQUE NOT NULL,
                nombre TEXT NOT NULL,
                categoria TEXT NOT NULL,
                categoria_id INTEGER,
                codigo_id INTEGER,
                cantidad INTEGER NOT NULL DEFAULT 1,
                estado TEXT NOT NULL,
                observaciones TEXT,
                seccion_id INTEGER NOT NULL,
                fecha_ingreso TEXT,
                FOREIGN KEY (seccion_id) REFERENCES secciones (id),
                FOREIGN KEY (categoria_id) REFERENCES categorias (id),
                FOREIGN KEY (codigo_id) REFERENCES codigos (id)
            )
        ''')

        # --- Migración simple: si la tabla 'elementos' ya existía sin estas
        # columnas (base de datos previa a esta versión), se agregan ahora. ---
        columnas_actuales = [c['name'] for c in db.execute("PRAGMA table_info(elementos)").fetchall()]
        if 'categoria_id' not in columnas_actuales:
            db.execute("ALTER TABLE elementos ADD COLUMN categoria_id INTEGER")
        if 'codigo_id' not in columnas_actuales:
            db.execute("ALTER TABLE elementos ADD COLUMN codigo_id INTEGER")
        if 'fecha_ingreso' not in columnas_actuales:
            db.execute("ALTER TABLE elementos ADD COLUMN fecha_ingreso TEXT")

        # Los elementos que ya existían antes de esta versión no tienen fecha.
        # Se les asigna una fecha "piso" para que siempre queden al final al
        # ordenar de más reciente a más antiguo (en vez de romper el orden).
        db.execute("UPDATE elementos SET fecha_ingreso = '2000-01-01 00:00:00' WHERE fecha_ingreso IS NULL")

        cursor = db.cursor()

        cursor.execute("SELECT COUNT(*) FROM secciones")
        if cursor.fetchone()[0] == 0:
            db.execute("INSERT INTO secciones (nombre, tipo) VALUES ('Laboratorio 1', 'Laboratorio')")
            db.execute("INSERT INTO secciones (nombre, tipo) VALUES ('Laboratorio 2', 'Laboratorio')")
            db.execute("INSERT INTO secciones (nombre, tipo) VALUES ('Taller 1', 'Taller')")

        cursor.execute("SELECT COUNT(*) FROM usuarios")
        if cursor.fetchone()[0] == 0:
            db.execute("INSERT INTO usuarios (username, password, rol) VALUES ('admin', 'admin123', 'Administrador')")
            db.execute("INSERT INTO usuarios (username, password, rol) VALUES ('auxiliar', 'aux123', 'Auxiliar')")
            db.execute("INSERT INTO usuarios (username, password, rol) VALUES ('auditor', 'aud123', 'Auditor')")

        # --- Datos de ejemplo para categorías y códigos (solo la primera vez) ---
        cursor.execute("SELECT COUNT(*) FROM categorias")
        if cursor.fetchone()[0] == 0:
            db.execute("INSERT INTO categorias (nombre) VALUES ('Cómputo')")
            db.execute("INSERT INTO categorias (nombre) VALUES ('Red')")
            db.execute("INSERT INTO categorias (nombre) VALUES ('Herramienta')")

            cat_computo = db.execute("SELECT id FROM categorias WHERE nombre = 'Cómputo'").fetchone()['id']
            db.execute("INSERT INTO codigos (categoria_id, prefijo, descripcion) VALUES (?, 'MON', 'Monitor')", (cat_computo,))
            db.execute("INSERT INTO codigos (categoria_id, prefijo, descripcion) VALUES (?, 'CPU', 'Unidad Central de Procesamiento')", (cat_computo,))
            db.execute("INSERT INTO codigos (categoria_id, prefijo, descripcion) VALUES (?, 'TEC', 'Teclado')", (cat_computo,))
            db.execute("INSERT INTO codigos (categoria_id, prefijo, descripcion) VALUES (?, 'MOU', 'Mouse')", (cat_computo,))

        db.commit()

# Decorador: exige que haya una sesión iniciada
def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'username' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function

# Decorador: exige uno de los roles indicados. Ej: @rol_required('Administrador')
def rol_required(*roles_permitidos):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if session.get('rol') not in roles_permitidos:
                return "Acceso denegado", 403
            return f(*args, **kwargs)
        return decorated_function
    return decorator

# --- AUTENTICACIÓN ---

@app.route('/login', methods=['GET', 'POST'])
def login():
    error = None
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']

        db = get_db()
        user = db.execute('SELECT * FROM usuarios WHERE username = ? AND password = ?', (username, password)).fetchone()

        if user:
            session['username'] = user['username']
            session['rol'] = user['rol']
            return redirect(url_for('dashboard'))
        else:
            error = 'Usuario o contraseña incorrectos.'

    return render_template('login.html', error=error)

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

# --- DASHBOARD PRINCIPAL ---

@app.route('/')
@login_required
def index():
    return redirect(url_for('dashboard'))

@app.route('/dashboard')
@login_required
def dashboard():
    db = get_db()
    secciones = db.execute("SELECT * FROM secciones").fetchall()

    resumen_secciones = []
    for sec in secciones:
        total = db.execute("SELECT SUM(cantidad) FROM elementos WHERE seccion_id = ?", (sec['id'],)).fetchone()[0] or 0
        operativos = db.execute("SELECT SUM(cantidad) FROM elementos WHERE seccion_id = ? AND estado = 'Operativo'", (sec['id'],)).fetchone()[0] or 0
        mantenimiento = db.execute("SELECT SUM(cantidad) FROM elementos WHERE seccion_id = ? AND estado = 'Mantenimiento'", (sec['id'],)).fetchone()[0] or 0
        inoperativos = db.execute("SELECT SUM(cantidad) FROM elementos WHERE seccion_id = ? AND estado = 'Inoperativo'", (sec['id'],)).fetchone()[0] or 0

        resumen_secciones.append({
            'id': sec['id'],
            'nombre': sec['nombre'],
            'tipo': sec['tipo'],
            'total_equipos': total,
            'operativos': operativos,
            'mantenimiento': mantenimiento,
            'inoperativos': inoperativos
        })

    ultimas_observaciones = db.execute('''
        SELECT e.*, s.nombre as seccion_nombre
        FROM elementos e
        JOIN secciones s ON e.seccion_id = s.id
        WHERE e.observaciones IS NOT NULL AND e.observaciones != ''
        ORDER BY e.id DESC LIMIT 5
    ''').fetchall()

    return render_template('dashboard.html',
                           secciones=secciones,
                           resumen_secciones=resumen_secciones,
                           ultimas_observaciones=ultimas_observaciones)

# --- INVENTARIO DE SECCIONES ---

@app.route('/seccion/<int:seccion_id>')
@login_required
def ver_seccion(seccion_id):
    db = get_db()
    secciones = db.execute("SELECT * FROM secciones").fetchall()
    seccion_actual = db.execute("SELECT * FROM secciones WHERE id = ?", (seccion_id,)).fetchone()
    elementos = db.execute(
        "SELECT * FROM elementos WHERE seccion_id = ? ORDER BY fecha_ingreso DESC, id DESC",
        (seccion_id,)
    ).fetchall()

    # Catálogo de categorías/códigos para los formularios de ingreso
    categorias = db.execute("SELECT * FROM categorias ORDER BY nombre").fetchall()
    codigos = db.execute("SELECT * FROM codigos ORDER BY prefijo").fetchall()

    # id -> prefijo oficial (para elementos que sí quedaron ligados a un código)
    prefijo_por_codigo_id = {c['id']: c['prefijo'] for c in codigos}
    # prefijo -> descripción (para mostrarla en el encabezado del acordeón)
    descripcion_por_prefijo = {c['prefijo']: (c['descripcion'] or '') for c in codigos}

    # --- Agrupar elementos por prefijo de código (MON, CPU, TEC...) para el acordeón ---
    grupos_dict = {}
    for elem in elementos:
        if elem['codigo_id'] and elem['codigo_id'] in prefijo_por_codigo_id:
            prefijo = prefijo_por_codigo_id[elem['codigo_id']]
        else:
            # Elementos antiguos sin codigo_id: se estima el prefijo desde el texto del código.
            codigo_texto = elem['codigo'] or ''
            prefijo = codigo_texto.split('-')[0] if '-' in codigo_texto else codigo_texto

        if prefijo not in grupos_dict:
            grupos_dict[prefijo] = {
                'prefijo': prefijo,
                'categoria': elem['categoria'],
                'descripcion': descripcion_por_prefijo.get(prefijo, ''),
                'elementos': []
            }
        grupos_dict[prefijo]['elementos'].append(elem)

    grupos = []
    for prefijo, g in grupos_dict.items():
        elementos_grupo = g['elementos']
        fechas = [i['fecha_ingreso'] for i in elementos_grupo if i['fecha_ingreso']]
        grupos.append({
            'prefijo': prefijo,
            'categoria': g['categoria'],
            'descripcion': g['descripcion'],
            'elementos': elementos_grupo,
            'cantidad_registros': len(elementos_grupo),
            'total_cantidad': sum(i['cantidad'] for i in elementos_grupo),
            'operativos': sum(i['cantidad'] for i in elementos_grupo if i['estado'] == 'Operativo'),
            'mantenimiento': sum(i['cantidad'] for i in elementos_grupo if i['estado'] == 'Mantenimiento'),
            'inoperativos': sum(i['cantidad'] for i in elementos_grupo if i['estado'] == 'Inoperativo'),
            'fecha_reciente': max(fechas) if fechas else ''
        })

    # Recién ingresados primero (coincide con el filtro "Recientes primero" por defecto)
    grupos.sort(key=lambda g: g['fecha_reciente'], reverse=True)

    categorias_presentes = sorted({g['categoria'] for g in grupos})
    prefijos_presentes = sorted({g['prefijo'] for g in grupos})

    return render_template('inventario.html',
                           secciones=secciones,
                           seccion_actual=seccion_actual,
                           grupos=grupos,
                           categorias=categorias,
                           codigos=codigos,
                           categorias_presentes=categorias_presentes,
                           prefijos_presentes=prefijos_presentes)

# --- GESTIÓN DE CATEGORÍAS Y CÓDIGOS (Solo Administrador) ---

@app.route('/categorias')
@login_required
@rol_required('Administrador')
def categorias():
    db = get_db()
    secciones = db.execute("SELECT * FROM secciones").fetchall()
    lista_categorias = db.execute("SELECT * FROM categorias ORDER BY nombre").fetchall()
    lista_codigos = db.execute('''
        SELECT cod.*, cat.nombre as categoria_nombre
        FROM codigos cod
        JOIN categorias cat ON cod.categoria_id = cat.id
        ORDER BY cat.nombre, cod.prefijo
    ''').fetchall()

    return render_template('categorias.html',
                           secciones=secciones,
                           categorias=lista_categorias,
                           codigos=lista_codigos)

@app.route('/categorias/agregar', methods=['POST'])
@login_required
@rol_required('Administrador')
def agregar_categoria():
    nombre = request.form['nombre'].strip()
    db = get_db()
    try:
        db.execute("INSERT INTO categorias (nombre) VALUES (?)", (nombre,))
        db.commit()
    except sqlite3.IntegrityError:
        print(f"Error: la categoría '{nombre}' ya existe.")
    return redirect(url_for('categorias'))

@app.route('/categorias/eliminar/<int:categoria_id>')
@login_required
@rol_required('Administrador')
def eliminar_categoria(categoria_id):
    db = get_db()
    # No se puede borrar una categoría que ya tiene códigos asociados,
    # para no dejar códigos "huérfanos" ni romper elementos existentes.
    en_uso = db.execute("SELECT COUNT(*) FROM codigos WHERE categoria_id = ?", (categoria_id,)).fetchone()[0]
    if en_uso == 0:
        db.execute("DELETE FROM categorias WHERE id = ?", (categoria_id,))
        db.commit()
    return redirect(url_for('categorias'))

@app.route('/codigos/agregar', methods=['POST'])
@login_required
@rol_required('Administrador')
def agregar_codigo():
    categoria_id = request.form['categoria_id']
    prefijo = request.form['prefijo'].strip().upper()
    descripcion = request.form.get('descripcion', '').strip()

    db = get_db()
    try:
        db.execute(
            "INSERT INTO codigos (categoria_id, prefijo, descripcion) VALUES (?, ?, ?)",
            (categoria_id, prefijo, descripcion)
        )
        db.commit()
    except sqlite3.IntegrityError:
        print(f"Error: el prefijo '{prefijo}' ya existe.")
    return redirect(url_for('categorias'))

@app.route('/codigos/eliminar/<int:codigo_id>')
@login_required
@rol_required('Administrador')
def eliminar_codigo(codigo_id):
    db = get_db()
    # No se puede borrar un código que ya se usó en algún elemento del inventario.
    en_uso = db.execute("SELECT COUNT(*) FROM elementos WHERE codigo_id = ?", (codigo_id,)).fetchone()[0]
    if en_uso == 0:
        db.execute("DELETE FROM codigos WHERE id = ?", (codigo_id,))
        db.commit()
    return redirect(url_for('categorias'))

# --- REORDENAMIENTO AUTOMÁTICO DE CÓDIGOS ---

def recalcular_codigos(db, codigo_id):
    """Renumera todos los elementos de un código (prefijo) para que queden
    consecutivos desde 001, respetando su orden actual, y actualiza el
    contador del prefijo al total real de elementos.

    Ejemplo: existen MON-001..MON-005 y se elimina MON-002
             -> MON-003 pasa a MON-002, MON-004 a MON-003, MON-005 a MON-004
             -> contador = 4, el siguiente nuevo será MON-005.
    """
    codigo_row = db.execute("SELECT prefijo FROM codigos WHERE id = ?", (codigo_id,)).fetchone()
    if not codigo_row:
        return
    prefijo = codigo_row['prefijo']

    elementos = db.execute(
        "SELECT id, codigo FROM elementos WHERE codigo_id = ?", (codigo_id,)
    ).fetchall()

    def numero(e):
        try:
            return int(e['codigo'].rsplit('-', 1)[1])
        except (IndexError, ValueError):
            return 0

    # Orden ascendente por número actual: cada elemento recibe un número
    # menor o igual al que tenía, así nunca choca con el UNIQUE de 'codigo'.
    elementos = sorted(elementos, key=numero)

    for posicion, e in enumerate(elementos, start=1):
        nuevo_codigo = f"{prefijo}-{posicion:03d}"
        if nuevo_codigo != e['codigo']:
            db.execute("UPDATE elementos SET codigo = ? WHERE id = ?", (nuevo_codigo, e['id']))

    db.execute("UPDATE codigos SET contador = ? WHERE id = ?", (len(elementos), codigo_id))

# --- OPERACIONES CRUD DE ELEMENTOS ---

@app.route('/agregar', methods=['POST'])
@login_required
@rol_required('Administrador', 'Auxiliar')
def agregar_elemento():
    codigo_id = request.form['codigo_id']
    nombre = request.form['nombre']
    estado = request.form['estado']
    observaciones = request.form['observaciones']
    seccion_id = request.form['seccion_id']

    db = get_db()

    codigo_row = db.execute("SELECT * FROM codigos WHERE id = ?", (codigo_id,)).fetchone()
    if not codigo_row:
        return "Código no válido", 400

    categoria_row = db.execute("SELECT * FROM categorias WHERE id = ?", (codigo_row['categoria_id'],)).fetchone()

    # --- Autonumeración: MON-001, MON-002, MON-003... ---
    nuevo_contador = codigo_row['contador'] + 1
    codigo_generado = f"{codigo_row['prefijo']}-{nuevo_contador:03d}"
    fecha_ingreso = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

    try:
        db.execute("UPDATE codigos SET contador = ? WHERE id = ?", (nuevo_contador, codigo_id))
        db.execute('''
            INSERT INTO elementos (codigo, nombre, categoria, categoria_id, codigo_id, cantidad, estado, observaciones, seccion_id, fecha_ingreso)
            VALUES (?, ?, ?, ?, ?, 1, ?, ?, ?, ?)
        ''', (codigo_generado, nombre, categoria_row['nombre'], categoria_row['id'], codigo_id,
              estado, observaciones, seccion_id, fecha_ingreso))
        db.commit()
    except sqlite3.IntegrityError:
        db.rollback()
        print("Error: no se pudo generar el elemento (código duplicado).")

    return redirect(url_for('ver_seccion', seccion_id=seccion_id))

@app.route('/agregar_lote', methods=['POST'])
@login_required
@rol_required('Administrador', 'Auxiliar')
def agregar_elemento_lote():
    # Ingreso masivo: crea N elementos individuales idénticos (mismo nombre,
    # categoría, estado inicial y observaciones), cada uno con su propio
    # código autonumerado y cantidad=1, para poder rastrearlos por separado
    # más adelante (ej. si uno se daña, no afecta el registro de los demás).
    codigo_id = request.form['codigo_id']
    nombre = request.form['nombre']
    estado = request.form['estado']
    observaciones = request.form['observaciones']
    seccion_id = request.form['seccion_id']

    try:
        cantidad_lote = int(request.form['cantidad_lote'])
    except (KeyError, ValueError):
        return "Cantidad de lote no válida", 400

    if cantidad_lote < 1:
        return "La cantidad del lote debe ser al menos 1", 400

    db = get_db()

    codigo_row = db.execute("SELECT * FROM codigos WHERE id = ?", (codigo_id,)).fetchone()
    if not codigo_row:
        return "Código no válido", 400

    categoria_row = db.execute("SELECT * FROM categorias WHERE id = ?", (codigo_row['categoria_id'],)).fetchone()
    fecha_ingreso = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    contador = codigo_row['contador']

    try:
        for _ in range(cantidad_lote):
            contador += 1
            codigo_generado = f"{codigo_row['prefijo']}-{contador:03d}"
            db.execute('''
                INSERT INTO elementos (codigo, nombre, categoria, categoria_id, codigo_id, cantidad, estado, observaciones, seccion_id, fecha_ingreso)
                VALUES (?, ?, ?, ?, ?, 1, ?, ?, ?, ?)
            ''', (codigo_generado, nombre, categoria_row['nombre'], categoria_row['id'], codigo_id,
                  estado, observaciones, seccion_id, fecha_ingreso))

        db.execute("UPDATE codigos SET contador = ? WHERE id = ?", (contador, codigo_id))
        db.commit()
    except sqlite3.IntegrityError:
        db.rollback()
        print("Error: no se pudo generar el lote (código duplicado).")

    return redirect(url_for('ver_seccion', seccion_id=seccion_id))

@app.route('/editar/<int:elem_id>', methods=['POST'])
@login_required
@rol_required('Administrador', 'Auxiliar')
def editar_elemento(elem_id):
    # La categoría, el código y la cantidad ya no se editan aquí: cada
    # elemento representa una sola unidad y su código queda fijo desde el ingreso.
    nombre = request.form['nombre']
    estado = request.form['estado']
    observaciones = request.form['observaciones']
    seccion_id = request.form['seccion_id']

    db = get_db()
    db.execute('''
        UPDATE elementos
        SET nombre = ?, estado = ?, observaciones = ?
        WHERE id = ?
    ''', (nombre, estado, observaciones, elem_id))
    db.commit()

    return redirect(url_for('ver_seccion', seccion_id=seccion_id))

@app.route('/eliminar/<int:elem_id>/<int:seccion_id>')
@login_required
@rol_required('Administrador')
def eliminar_elemento(elem_id, seccion_id):
    db = get_db()

    # Se guarda a qué código (MON, CPU...) pertenecía antes de borrarlo
    fila = db.execute('SELECT codigo_id FROM elementos WHERE id = ?', (elem_id,)).fetchone()
    codigo_id = fila['codigo_id'] if fila else None

    db.execute('DELETE FROM elementos WHERE id = ?', (elem_id,))

    # Reordena la numeración: sin importar si se borró el primero, uno del
    # medio o el último, los códigos restantes quedan consecutivos (001, 002...)
    # y el contador vuelve a coincidir con la cantidad real de elementos.
    if codigo_id:
        recalcular_codigos(db, codigo_id)

    db.commit()
    return redirect(url_for('ver_seccion', seccion_id=seccion_id))

# --- EXPORTACIÓN DE DATOS (Solo Administrador) ---

def obtener_datos_exportacion():
    """Recolecta el resumen por sección y el detalle de todos los elementos
    de todas las secciones, para usarlos tanto en el Excel como en el PDF."""
    db = get_db()
    secciones = db.execute("SELECT * FROM secciones").fetchall()

    resumen = []
    detalle = []

    for sec in secciones:
        elementos_sec = db.execute(
            "SELECT * FROM elementos WHERE seccion_id = ? ORDER BY codigo",
            (sec['id'],)
        ).fetchall()

        operativos = sum(e['cantidad'] for e in elementos_sec if e['estado'] == 'Operativo')
        mantenimiento = sum(e['cantidad'] for e in elementos_sec if e['estado'] == 'Mantenimiento')
        inoperativos = sum(e['cantidad'] for e in elementos_sec if e['estado'] == 'Inoperativo')

        resumen.append({
            'nombre': sec['nombre'],
            'tipo': sec['tipo'],
            'total': operativos + mantenimiento + inoperativos,
            'operativos': operativos,
            'mantenimiento': mantenimiento,
            'inoperativos': inoperativos
        })

        for e in elementos_sec:
            detalle.append({
                'seccion': sec['nombre'],
                'codigo': e['codigo'],
                'nombre': e['nombre'],
                'categoria': e['categoria'],
                'estado': e['estado'],
                'fecha_ingreso': e['fecha_ingreso'][:10] if e['fecha_ingreso'] else '-',
                'observaciones': e['observaciones'] or ''
            })

    return resumen, detalle

def generar_excel_inventario():
    resumen, detalle = obtener_datos_exportacion()
    color_encabezado = '6750A4'

    def pintar_encabezado(ws, encabezados):
        ws.append(encabezados)
        for col in range(1, len(encabezados) + 1):
            celda = ws.cell(row=1, column=col)
            celda.font = Font(bold=True, color='FFFFFF')
            celda.fill = PatternFill('solid', fgColor=color_encabezado)
            celda.alignment = Alignment(horizontal='center')

    wb = Workbook()

    # --- Hoja 1: Resumen por sección ---
    ws_resumen = wb.active
    ws_resumen.title = 'Resumen'
    pintar_encabezado(ws_resumen, ['Sección', 'Tipo', 'Total', 'Operativos', 'Mantenimiento', 'Inoperativos'])
    for fila in resumen:
        ws_resumen.append([
            fila['nombre'], fila['tipo'], fila['total'],
            fila['operativos'], fila['mantenimiento'], fila['inoperativos']
        ])
    for col, ancho in enumerate([22, 16, 10, 12, 14, 12], start=1):
        ws_resumen.column_dimensions[get_column_letter(col)].width = ancho

    # --- Hoja 2: Detalle de todos los elementos, de todas las secciones ---
    ws_detalle = wb.create_sheet('Inventario Detallado')
    pintar_encabezado(ws_detalle, ['Sección', 'Código', 'Elemento', 'Categoría', 'Estado', 'Fecha de Ingreso', 'Observaciones'])
    for fila in detalle:
        ws_detalle.append([
            fila['seccion'], fila['codigo'], fila['nombre'], fila['categoria'],
            fila['estado'], fila['fecha_ingreso'], fila['observaciones']
        ])
    for col, ancho in enumerate([22, 14, 30, 18, 16, 18, 45], start=1):
        ws_detalle.column_dimensions[get_column_letter(col)].width = ancho

    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida

def generar_pdf_inventario():
    resumen, detalle = obtener_datos_exportacion()
    estilos = getSampleStyleSheet()
    color_encabezado = colors.HexColor('#6750A4')
    color_fila_alterna = colors.HexColor('#F4F6F9')

    salida = BytesIO()
    doc = SimpleDocTemplate(
        salida, pagesize=landscape(letter),
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm
    )

    cuerpo = [
        Paragraph('Reporte de Inventario — Control Inventory', estilos['Title']),
        Paragraph(f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}", estilos['Normal']),
        Spacer(1, 0.6 * cm),
        Paragraph('Resumen por Sección', estilos['Heading2']),
    ]

    datos_resumen = [['Sección', 'Tipo', 'Total', 'Operativos', 'Mantenimiento', 'Inoperativos']]
    for fila in resumen:
        datos_resumen.append([
            fila['nombre'], fila['tipo'], str(fila['total']),
            str(fila['operativos']), str(fila['mantenimiento']), str(fila['inoperativos'])
        ])
    tabla_resumen = Table(datos_resumen, repeatRows=1)
    tabla_resumen.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), color_encabezado),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, color_fila_alterna]),
    ]))
    cuerpo.append(tabla_resumen)
    cuerpo.append(Spacer(1, 1 * cm))

    cuerpo.append(Paragraph('Inventario Detallado', estilos['Heading2']))
    datos_detalle = [['Sección', 'Código', 'Elemento', 'Categoría', 'Estado', 'Ingreso', 'Observaciones']]
    for fila in detalle:
        datos_detalle.append([
            fila['seccion'], fila['codigo'], fila['nombre'], fila['categoria'],
            fila['estado'], fila['fecha_ingreso'], fila['observaciones'][:70]
        ])
    tabla_detalle = Table(
        datos_detalle, repeatRows=1,
        colWidths=[3 * cm, 2.2 * cm, 4.5 * cm, 2.8 * cm, 2.5 * cm, 2.3 * cm, 6.5 * cm]
    )
    tabla_detalle.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), color_encabezado),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, color_fila_alterna]),
    ]))
    cuerpo.append(tabla_detalle)

    doc.build(cuerpo)
    salida.seek(0)
    return salida

@app.route('/exportar/excel')
@login_required
@rol_required('Administrador')
def exportar_excel():
    archivo = generar_excel_inventario()
    return send_file(
        archivo,
        as_attachment=True,
        download_name=f"inventario_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@app.route('/exportar/pdf')
@login_required
@rol_required('Administrador')
def exportar_pdf():
    archivo = generar_pdf_inventario()
    return send_file(
        archivo,
        as_attachment=True,
        download_name=f"inventario_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf",
        mimetype='application/pdf'
    )

# Se ejecuta también bajo gunicorn (donde __main__ no se cumple)
init_db()

if __name__ == '__main__':
    port = int(os.environ.get("PORT", 5000))
    app.run(host='0.0.0.0', port=port, debug=True)